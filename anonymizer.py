import re
from typing import Callable, Dict, Iterable

from docx import Document
from docx.text.paragraph import Paragraph


FullGenerator = Callable[[str, str, str], str]
InitGenerator = Callable[[str, str], str]


def merge_runs_in_paragraph(paragraph: Paragraph) -> None:
    full_text = "".join(run.text for run in paragraph.runs)
    if len(paragraph.runs) <= 1:
        return
    first = paragraph.runs[0]
    style = first.style
    for run in paragraph.runs:
        run.text = ""
    new_run = paragraph.add_run(full_text)
    try:
        new_run.style = style
    except Exception:
        pass


def replace_full_pib_in_paragraph(
    paragraph: Paragraph,
    pattern_full: re.Pattern,
    replace_map_full: Dict[str, str],
    generator_full: FullGenerator,
) -> None:
    text = paragraph.text

    def repl_full(match: re.Match) -> str:
        old_surname = match.group(1)
        old_name = match.group(2)
        old_patronymic = match.group(3)
        key = f"{old_surname} {old_name} {old_patronymic}"
        if key in replace_map_full:
            return replace_map_full[key]
        new_full = generator_full(old_surname, old_name, old_patronymic)
        replace_map_full[key] = new_full
        return new_full

    new_text = pattern_full.sub(repl_full, text)
    if new_text != text:
        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = new_text
        else:
            paragraph.add_run(new_text)


def replace_initials_in_paragraph(
    paragraph: Paragraph,
    pattern_init: re.Pattern,
    replace_map_init: Dict[str, str],
    generator_init: InitGenerator,
) -> None:
    text = paragraph.text

    def repl_init(match: re.Match) -> str:
        old_surname = match.group(1)
        old_initial = match.group(2)
        key = f"{old_surname} {old_initial}"
        if key in replace_map_init:
            return replace_map_init[key]
        new_init = generator_init(old_surname, old_initial)
        replace_map_init[key] = new_init
        return new_init

    new_text = pattern_init.sub(repl_init, text)
    if new_text != text:
        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = new_text
        else:
            paragraph.add_run(new_text)


def process_paragraphs(
    paragraphs: Iterable[Paragraph],
    pattern_full: re.Pattern,
    pattern_init: re.Pattern,
    replace_map_full: Dict[str, str],
    replace_map_init: Dict[str, str],
    generator_full: FullGenerator,
    generator_init: InitGenerator,
    replace_full: bool,
    replace_initials: bool,
) -> None:
    for paragraph in paragraphs:
        merge_runs_in_paragraph(paragraph)
        if replace_full:
            replace_full_pib_in_paragraph(
                paragraph,
                pattern_full,
                replace_map_full,
                generator_full,
            )
        if replace_initials:
            replace_initials_in_paragraph(
                paragraph,
                pattern_init,
                replace_map_init,
                generator_init,
            )


def process_docx(
    input_path: str,
    output_path: str,
    generator_full: FullGenerator,
    generator_init: InitGenerator,
    replace_full: bool = True,
    replace_initials: bool = True,
    process_tables: bool = True,
    process_headers_footers: bool = True,
) -> None:
    pattern_full = re.compile(
        r"\b([А-ЯҐЄІЇ]+)\s+([А-ЯҐЄІЇ][а-яґєії]+)\s+([А-ЯҐЄІЇ][а-яґєії]+)\b"
    )
    pattern_init = re.compile(
        r"\b([А-ЯҐЄІЇ][а-яґєії]+)\s+([А-ЯҐЄІЇ])\.\s"
    )

    doc = Document(input_path)
    replace_map_full: Dict[str, str] = {}
    replace_map_init: Dict[str, str] = {}

    process_paragraphs(
        doc.paragraphs,
        pattern_full,
        pattern_init,
        replace_map_full,
        replace_map_init,
        generator_full,
        generator_init,
        replace_full,
        replace_initials,
    )

    if process_tables:
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    process_paragraphs(
                        cell.paragraphs,
                        pattern_full,
                        pattern_init,
                        replace_map_full,
                        replace_map_init,
                        generator_full,
                        generator_init,
                        replace_full,
                        replace_initials,
                    )

    if process_headers_footers:
        for section in doc.sections:
            process_paragraphs(
                section.header.paragraphs,
                pattern_full,
                pattern_init,
                replace_map_full,
                replace_map_init,
                generator_full,
                generator_init,
                replace_full,
                replace_initials,
            )
            process_paragraphs(
                section.footer.paragraphs,
                pattern_full,
                pattern_init,
                replace_map_full,
                replace_map_init,
                generator_full,
                generator_init,
                replace_full,
                replace_initials,
            )
            if process_tables:
                for table in section.header.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            process_paragraphs(
                                cell.paragraphs,
                                pattern_full,
                                pattern_init,
                                replace_map_full,
                                replace_map_init,
                                generator_full,
                                generator_init,
                                replace_full,
                                replace_initials,
                            )
                for table in section.footer.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            process_paragraphs(
                                cell.paragraphs,
                                pattern_full,
                                pattern_init,
                                replace_map_full,
                                replace_map_init,
                                generator_full,
                                generator_init,
                                replace_full,
                                replace_initials,
                            )

    doc.save(output_path)
